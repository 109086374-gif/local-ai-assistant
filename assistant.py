#!/usr/bin/env python3
"""
assistant.py -- 本地 CLI 协调器（调用 whisper.cpp / llama.cpp / python 工具）
用法示例：
  python assistant.py transcribe /path/to/input.wav
  python assistant.py write "写一封邮件，主题：迟到说明"
  python assistant.py genppt "报告主题" output.pptx
  python assistant.py analyze sample.csv
  python assistant.py codepack myscript.py
注意：需要先把 whisper / llama 可执行文件放到 bin/，并把模型放到 models/
"""
import sys
import subprocess
import os
from pathlib import Path

ROOT = Path("D:/assistant")
BIN = ROOT / "bin"
MODELS = ROOT / "models"

def run_whisper(wav_path):
    model = MODELS / "ggml-base.bin"
    exe = BIN / "whisper"
    if not exe.exists():
        print("找不到 whisper 可执行文件：", exe)
        return
    if not model.exists():
        print("找不到模型，请把 ggml-base.bin 放到", MODELS)
        return
    cmd = [str(exe), "-m", str(model), "-f", str(wav_path), "-l", "zh", "-otxt"]
    subprocess.run(cmd)

def run_llama_prompt(prompt):
    model = MODELS / "ggml-7b-q4.bin"
    exe = BIN / "llama"
    if not exe.exists():
        print("找不到 llama 可执行文件：", exe)
        return
    if not model.exists():
        print("找不到 LLM 模型，请把量化 7B 放到", MODELS)
        return
    cmd = [str(exe), "-m", str(model), "-p", prompt]
    subprocess.run(cmd)

def gen_ppt(title, outpath):
    # 简单示例：把标题写到一页 PPT，用 python-pptx 可扩展
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_tf = slide.shapes.title
    title_tf.text = title
    prs.save(outpath)
    print("已保存", outpath)

def analyze_csv(path):
    import pandas as pd
    df = pd.read_csv(path)
    print("数据预览：")
    print(df.head())
    print("\n描述性统计：")
    print(df.describe())

def code_pack(pyfile):
    # 需要 pyinstaller 已安装（venv pip install pyinstaller）
    subprocess.run(["pyinstaller", "--onefile", pyfile])

def main():
    if len(sys.argv) < 2:
        print("用法: python assistant.py [transcribe|write|genppt|analyze|codepack] args...")
        return
    cmd = sys.argv[1]
    if cmd == "transcribe":
        if len(sys.argv) < 3:
            print("用法: assistant.py transcribe path.wav")
            return
        run_whisper(sys.argv[2])
    elif cmd == "write":
        prompt = " ".join(sys.argv[2:])
        run_llama_prompt(prompt)
    elif cmd == "genppt":
        title = sys.argv[2] if len(sys.argv) > 2 else "演示"
        out = sys.argv[3] if len(sys.argv) > 3 else "out.pptx"
        gen_ppt(title, out)
    elif cmd == "analyze":
        analyze_csv(sys.argv[2])
    elif cmd == "codepack":
        code_pack(sys.argv[2])
    else:
        print("未知命令", cmd)

if __name__ == "__main__":
    main()
